import pandas as pd
import os
import glob
import logging
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font, PatternFill

def consolidate_retention_reports(reports_dir="reports", output_file="reports/consolidado_retencion.xlsx"):
    """
    Consolida todos los archivos de retención en un único archivo Excel con formato de tabla.
    
    Args:
        reports_dir (str): Directorio donde están los reportes de retención
        output_file (str): Ruta del archivo consolidado de salida
    """
    try:
        # Configurar logging
        logging.basicConfig(level=logging.INFO)
        
        # Buscar todos los archivos de retención
        pattern = os.path.join(reports_dir, "retencion_periodo_*.xlsx")
        retention_files = glob.glob(pattern)
        
        if not retention_files:
            logging.warning("No se encontraron archivos de retención para consolidar.")
            return
        
        # Lista para almacenar todos los DataFrames
        all_data = []
        
        # Leer cada archivo y agregarlo a la lista
        for file in retention_files:
            try:
                # Extraer el periodo del nombre del archivo
                periodo = file.split("_")[-1].replace(".xlsx", "")
                
                # Leer el archivo
                df = pd.read_excel(file)
                
                # Agregar columna de fecha de actualización
                df['Fecha_Actualizacion'] = pd.Timestamp.now().strftime('%Y-%m-%d %H:%M:%S')
                
                # Agregar a la lista
                all_data.append(df)
                
                logging.info(f"Archivo procesado: {file}")
                
            except Exception as e:
                logging.error(f"Error procesando {file}: {str(e)}")
                continue
        
        if not all_data:
            logging.error("No se pudo procesar ningún archivo.")
            return
        
        # Consolidar todos los DataFrames
        consolidated_df = pd.concat(all_data, ignore_index=True)
        
        # Crear directorio si no existe
        os.makedirs(os.path.dirname(output_file), exist_ok=True)
        
        # Guardar como Excel con formato de tabla
        with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
            # Escribir los datos
            consolidated_df.to_excel(writer, sheet_name='Retencion', index=False)
            
            # Obtener la hoja de trabajo
            worksheet = writer.sheets['Retencion']
            
            # Obtener el rango de la tabla
            end_column = chr(65 + len(consolidated_df.columns) - 1)  # Última columna (A, B, C, etc.)
            end_row = len(consolidated_df) + 1  # +1 para incluir el encabezado
            table_range = f"A1:{end_column}{end_row}"
            
            # Crear y agregar la tabla
            table = Table(displayName="Retencion", ref=table_range)
            
            # Aplicar estilo a la tabla
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False
            )
            
            # Agregar la tabla a la hoja
            worksheet.add_table(table)
            
            # Dar formato a los encabezados
            header_font = Font(bold=True)
            header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
            
            for cell in worksheet[1]:
                cell.font = header_font
                cell.fill = header_fill
            
            # Ajustar el ancho de las columnas
            for idx, column in enumerate(worksheet.columns):
                max_length = 0
                column = [cell for cell in column]
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = (max_length + 2)
                worksheet.column_dimensions[column[0].column_letter].width = adjusted_width
        
        logging.info(f"Archivo consolidado generado exitosamente: {output_file}")
        logging.info(f"Total de registros consolidados: {len(consolidated_df)}")
        
    except Exception as e:
        logging.error(f"Error durante la consolidación: {str(e)}")
        raise

if __name__ == "__main__":
    consolidate_retention_reports()
